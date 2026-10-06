from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl

from src.domain.normalizers import cell_to_text, normalize_money
from src.parsers.xlsx_parser import XlsxPaymentParser


def test_xlsx_uses_read_only_data_only_and_preserves_visible_zero_format(
    tmp_path: Path,
    monkeypatch,
) -> None:
    path = tmp_path / "pagos.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append([" FECHA ", "MONTO", "RECAUDO", "DOCUMENTO", "OPERACION"])
    worksheet.append(["2026-10-01", 200, "BCP", 3241, 60814])
    worksheet["D2"].number_format = "00000000"
    worksheet["E2"].number_format = "0000000"
    workbook.save(path)
    original_loader = openpyxl.load_workbook
    captured: dict[str, object] = {}

    def recording_loader(*args, **kwargs):
        captured.update(kwargs)
        return original_loader(*args, **kwargs)

    monkeypatch.setattr("src.parsers.xlsx_parser.load_workbook", recording_loader)

    rows = XlsxPaymentParser().rows(path)
    row = next(rows)

    assert iter(rows) is rows
    assert captured["read_only"] is True
    assert captured["data_only"] is True
    assert cell_to_text(row.values["DOCUMENTO"]) == "00003241"
    assert cell_to_text(row.values["OPERACION"]) == "0060814"


def test_xlsx_numeric_money_cells_are_marked_for_binary_float_tolerance(tmp_path: Path) -> None:
    path = tmp_path / "pagos-float.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["FECHA", "MONTO", "ACUMULADO", "RECAUDO"])
    worksheet.append(
        ["2026-10-01", 216791.33000000002, 1000.1200000000001, "BCP"]
    )
    workbook.save(path)

    row = next(XlsxPaymentParser().rows(path))

    assert row.values["MONTO"].is_xlsx_float is True
    assert row.values["ACUMULADO"].is_xlsx_float is True
    assert normalize_money(row.values["MONTO"]) == Decimal("216791.33")
    assert normalize_money(row.values["ACUMULADO"], field_label="Acumulado") == Decimal(
        "1000.12"
    )
