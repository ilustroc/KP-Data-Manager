from __future__ import annotations

from pathlib import Path
from typing import Iterator
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..domain.errors import StructuralFileError
from ..domain.models import CellValue, SourceRow
from .base import EXPECTED_COLUMNS, build_header_map


class XlsxPaymentParser:
    def rows(self, path: Path) -> Iterator[SourceRow]:
        try:
            workbook = load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exception:
            raise StructuralFileError(f"XLSX inválido: {exception}.") from exception

        try:
            if not workbook.worksheets:
                raise StructuralFileError("El XLSX no contiene hojas.")
            worksheet = workbook.worksheets[0]
            row_iterator = worksheet.iter_rows()
            try:
                header_cells = next(row_iterator)
            except StopIteration as exception:
                raise StructuralFileError("El XLSX está vacío.") from exception
            header_map = build_header_map([cell.value for cell in header_cells])

            for cells in row_iterator:
                source_values = {
                    column: CellValue(
                        cells[index].value if index < len(cells) else None,
                        cells[index].number_format if index < len(cells) else None,
                    )
                    for column, index in header_map.items()
                }
                for optional_column in EXPECTED_COLUMNS:
                    source_values.setdefault(optional_column, CellValue(None))
                if all(cell.value is None or str(cell.value).strip() == "" for cell in source_values.values()):
                    continue
                yield SourceRow(number=cells[0].row, values=source_values)
        finally:
            workbook.close()
