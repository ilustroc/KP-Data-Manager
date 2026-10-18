from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator, Mapping, Sequence
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..domain.errors import StructuralFileError
from ..domain.models import CellValue, SearchSourceRow
class SearchParser:
    def __init__(
        self,
        extension: str,
        mapping: Sequence[Mapping[str, Any]],
        target_fields: Sequence[str],
    ) -> None:
        self.extension = extension
        self.mapping = tuple(mapping)
        self.target_fields = tuple(target_fields)

    def rows(self, path: Path) -> Iterator[SearchSourceRow]:
        if self.extension == "xlsx":
            yield from self._xlsx_rows(path)
            return
        if self.extension == "csv":
            yield from self._csv_rows(path)
            return
        raise StructuralFileError(f"Extensión SEARCH no soportada: {self.extension}.")

    def _xlsx_rows(self, path: Path) -> Iterator[SearchSourceRow]:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exception:
            raise StructuralFileError(f"XLSX inválido: {exception}.") from exception
        try:
            if not workbook.worksheets:
                raise StructuralFileError("El XLSX no contiene hojas.")
            worksheet = workbook.worksheets[0]
            iterator = worksheet.iter_rows()
            try:
                header_cells = next(iterator)
            except StopIteration as exception:
                raise StructuralFileError("El XLSX está vacío.") from exception
            headers = self._headers([cell.value for cell in header_cells])
            target_positions = self._validate_mapping(headers)
            for cells in iterator:
                values = tuple(
                    CellValue(
                        value=cell.value,
                        number_format=cell.number_format,
                        is_xlsx_float=isinstance(cell.value, float),
                        workbook_epoch=workbook.epoch,
                    )
                    for cell in cells[: len(headers)]
                )
                values += (CellValue(None),) * (len(headers) - len(values))
                if self._is_empty(values):
                    continue
                mapped = {target: values[position] for target, position in target_positions.items()}
                row_number = cells[0].row if cells else worksheet.max_row
                yield SearchSourceRow(row_number, headers, values, mapped)
        finally:
            workbook.close()

    def _csv_rows(self, path: Path) -> Iterator[SearchSourceRow]:
        try:
            with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as handle:
                sample = handle.read(32768)
                handle.seek(0)
                if not sample:
                    raise StructuralFileError("El CSV está vacío.")
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=";,").delimiter
                except csv.Error as exception:
                    raise StructuralFileError("No se pudo detectar el delimitador CSV ';' o ','.") from exception
                reader = csv.reader(handle, delimiter=delimiter)
                try:
                    headers = self._headers(next(reader))
                except StopIteration as exception:
                    raise StructuralFileError("El CSV no contiene encabezados.") from exception
                target_positions = self._validate_mapping(headers)
                for source_values in reader:
                    values = tuple(
                        CellValue(source_values[index] if index < len(source_values) else None)
                        for index in range(len(headers))
                    )
                    if self._is_empty(values):
                        continue
                    mapped = {target: values[position] for target, position in target_positions.items()}
                    yield SearchSourceRow(reader.line_num, headers, values, mapped)
        except UnicodeDecodeError as exception:
            raise StructuralFileError(
                f"El CSV no es UTF-8 válido cerca del byte {exception.start}."
            ) from exception

    @staticmethod
    def _headers(values: Sequence[Any]) -> tuple[str, ...]:
        headers = ["" if value is None else str(value).replace("\u00a0", " ").strip() for value in values]
        while headers and headers[-1] == "":
            headers.pop()
        if not headers or any(header == "" for header in headers):
            raise StructuralFileError("Todos los encabezados SEARCH deben tener nombre.")
        if len(set(headers)) != len(headers):
            raise StructuralFileError("El archivo SEARCH contiene encabezados duplicados.")
        return tuple(headers)

    def _validate_mapping(self, headers: tuple[str, ...]) -> dict[str, int]:
        mapped_targets: dict[str, int] = {}
        for row in self.mapping:
            target = str(row["target_column"])
            source = str(row["source_column"])
            position = int(row["source_position"]) - 1
            if target in mapped_targets or target not in self.target_fields:
                raise StructuralFileError("El mapping SEARCH contiene campos canónicos inválidos o duplicados.")
            if position < 0 or position >= len(headers) or headers[position] != source:
                raise StructuralFileError(
                    f"La columna mapeada '{source}' no coincide con la posición {position + 1}."
                )
            mapped_targets[target] = position
        missing = [field for field in self.target_fields if field not in mapped_targets]
        if missing:
            raise StructuralFileError("El mapping SEARCH está incompleto: " + ", ".join(missing))
        return mapped_targets

    @staticmethod
    def _is_empty(values: Sequence[CellValue]) -> bool:
        return all(cell.value is None or str(cell.value).strip() == "" for cell in values)
