from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterator
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..domain.errors import StructuralFileError
from ..domain.models import CellValue, SourceRow
from ..profiles.carteras.base import BaseCarteraProfile


class CsvCarteraParser:
    SAMPLE_LINES = 20
    SAMPLE_LINE_LIMIT = 32768

    def __init__(self, profile: BaseCarteraProfile) -> None:
        self.profile = profile

    def rows(self, path: Path) -> Iterator[SourceRow]:
        try:
            with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as handle:
                sample_parts: list[str] = []
                for _ in range(self.SAMPLE_LINES):
                    line = handle.readline(self.SAMPLE_LINE_LIMIT)
                    if line == "":
                        break
                    sample_parts.append(line)
                sample = "".join(sample_parts)
                handle.seek(0)
                if not sample:
                    raise StructuralFileError("El CSV está vacío.")
                delimiter = self._detect_delimiter(sample)
                reader = csv.reader(handle, delimiter=delimiter)
                try:
                    headers = next(reader)
                except StopIteration as exception:
                    raise StructuralFileError("El CSV no contiene encabezados.") from exception
                header_map = self.profile.build_header_map(headers)

                for values in reader:
                    source_values = {
                        column: CellValue(values[index] if index < len(values) else None)
                        for column, index in header_map.items()
                    }
                    for column in self.profile.expected_headers:
                        source_values.setdefault(column, CellValue(None))
                    if all(
                        cell.value is None or str(cell.value).strip() == ""
                        for cell in source_values.values()
                    ):
                        continue
                    yield SourceRow(number=reader.line_num, values=source_values)
        except UnicodeDecodeError as exception:
            raise StructuralFileError(
                f"El CSV no es UTF-8 válido cerca del byte {exception.start}."
            ) from exception
        except csv.Error as exception:
            raise StructuralFileError(f"CSV inválido: {exception}.") from exception

    @staticmethod
    def _detect_delimiter(sample: str) -> str:
        try:
            return csv.Sniffer().sniff(sample, delimiters=";,").delimiter
        except csv.Error as exception:
            raise StructuralFileError("No se pudo detectar el delimitador CSV ';' o ','.") from exception


class XlsxCarteraParser:
    def __init__(self, profile: BaseCarteraProfile) -> None:
        self.profile = profile

    def rows(self, path: Path) -> Iterator[SourceRow]:
        try:
            workbook = load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exception:
            raise StructuralFileError(f"XLSX inválido: {exception}.") from exception

        try:
            if not workbook.worksheets:
                raise StructuralFileError("El XLSX no contiene hojas.")
            worksheet = workbook.worksheets[0]
            row_iterator = worksheet.iter_rows()
            try:
                header_cells = next(row_iterator)
            except StopIteration as exception:
                raise StructuralFileError("El XLSX está vacío.") from exception
            header_map = self.profile.build_header_map([cell.value for cell in header_cells])

            for cells in row_iterator:
                source_values = {
                    column: self._cell_value(cells[index]) if index < len(cells) else CellValue(None)
                    for column, index in header_map.items()
                }
                for column in self.profile.expected_headers:
                    source_values.setdefault(column, CellValue(None))
                if all(
                    cell.value is None or str(cell.value).strip() == ""
                    for cell in source_values.values()
                ):
                    continue
                row_number = cells[0].row if cells else worksheet.max_row
                yield SourceRow(number=row_number, values=source_values)
        finally:
            workbook.close()

    @staticmethod
    def _cell_value(cell: Any) -> CellValue:
        return CellValue(
            value=cell.value,
            number_format=cell.number_format,
            is_xlsx_float=isinstance(cell.value, float),
        )
